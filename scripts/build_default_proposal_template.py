import os

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

OUTPUT_DIR = os.path.join("app", "proposal_templates")
OUTPUT_FILE = os.path.join(OUTPUT_DIR, "default_proposal.docx")

NAVY = RGBColor(0x1B, 0x24, 0x30)
COPPER = RGBColor(0xC9, 0x7B, 0x3D)
GREY = RGBColor(0x8A, 0x93, 0xA6)
BODY = RGBColor(0x2A, 0x33, 0x40)


def shade(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def cell_borders(table, color="E2DDD0", size="4"):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tbl_pr.append(borders)


def para(container, text, size=10, bold=False, color=BODY, align=None, space_after=4):
    p = container.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    if align is not None:
        p.alignment = align
    return p


def set_cell_text(cell, text, size=9, bold=False, color=BODY, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(2)
    if align is not None:
        p.alignment = align


def bottom_border(paragraph, size, color):
    p_pr = paragraph._p.get_or_add_pPr()
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), color)
    border.append(bottom)
    p_pr.append(border)


def section_title(doc, text):
    p = para(doc, text.upper(), size=8, bold=True, color=GREY, space_after=6)
    p.paragraph_format.space_before = Pt(14)
    bottom_border(p, "6", "E2DDD0")


def header_row(table, labels, widths, right_cols=()):
    row = table.rows[0]
    for i, label in enumerate(labels):
        cell = row.cells[i]
        cell.width = widths[i]
        shade(cell, "1B2430")
        set_cell_text(
            cell,
            label.upper(),
            size=8,
            bold=True,
            color=RGBColor(0xF6, 0xF3, 0xEC),
            align=WD_ALIGN_PARAGRAPH.RIGHT if i in right_cols else None,
        )


def build():
    doc = Document()

    sec = doc.sections[0]
    sec.page_width = Mm(210)
    sec.page_height = Mm(297)
    sec.left_margin = Mm(16)
    sec.right_margin = Mm(16)
    sec.top_margin = Mm(18)
    sec.bottom_margin = Mm(20)

    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(10)

    # Letterhead
    head = doc.add_table(rows=1, cols=2)
    head.alignment = WD_TABLE_ALIGNMENT.CENTER
    left, right = head.rows[0].cells
    left.width = Mm(110)
    right.width = Mm(66)

    left.text = ""
    p = left.paragraphs[0]
    r = p.add_run("{{company_name}}")
    r.font.name = "Arial"
    r.font.size = Pt(20)
    r.bold = True
    r.font.color.rgb = NAVY
    p2 = left.add_paragraph()
    r2 = p2.add_run("{{company_email}}")
    r2.font.name = "Arial"
    r2.font.size = Pt(8.5)
    r2.font.color.rgb = GREY

    right.text = ""
    rp = right.paragraphs[0]
    rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    rr = rp.add_run("{{quotation_number}}")
    rr.font.name = "Courier New"
    rr.font.size = Pt(13)
    rr.bold = True
    rr.font.color.rgb = COPPER
    for text in ("Version {{quotation_version}}", "{{quotation_date}}", "Valid until {{validity_date}}"):
        q = right.add_paragraph()
        q.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        qr = q.add_run(text)
        qr.font.name = "Arial"
        qr.font.size = Pt(9)
        qr.font.color.rgb = GREY

    rule = doc.add_paragraph()
    rule.paragraph_format.space_after = Pt(8)
    bottom_border(rule, "18", "C97B3D")

    # Title
    para(doc, "PROPOSAL & QUOTATION", size=16, bold=True, color=NAVY, space_after=2)
    para(doc, "Windows & Doors", size=10, color=GREY, space_after=10)

    # Customer block
    cust = doc.add_table(rows=1, cols=1)
    cust.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = cust.rows[0].cells[0]
    shade(c, "F6F3EC")
    c.text = ""
    lp = c.paragraphs[0]
    lr = lp.add_run("PREPARED FOR")
    lr.font.name = "Arial"
    lr.font.size = Pt(7.5)
    lr.bold = True
    lr.font.color.rgb = GREY
    np_ = c.add_paragraph()
    nr = np_.add_run("{{customer_name}}")
    nr.font.name = "Arial"
    nr.font.size = Pt(12)
    nr.bold = True
    nr.font.color.rgb = NAVY
    ap = c.add_paragraph()
    ar = ap.add_run("{{project_address}}")
    ar.font.name = "Arial"
    ar.font.size = Pt(9)
    ar.font.color.rgb = BODY

    # Executive summary
    section_title(doc, "Executive Summary")
    para(doc, "{{executive_summary}}", size=10)

    # Scope of work
    section_title(doc, "Scope of Work")
    para(doc, "{{scope_of_work}}", size=10)

    # Line items
    section_title(doc, "Specification & Pricing")
    widths = [Mm(10), Mm(58), Mm(30), Mm(32), Mm(12), Mm(34)]
    items = doc.add_table(rows=2, cols=6)
    items.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_borders(items)
    header_row(items, ["No", "Item", "Material", "Size", "Qty", "Amount"], widths, right_cols=(4, 5))
    tokens = [
        "{{item_no}}", "{{item_label}}", "{{item_material}}",
        "{{item_size}}", "{{item_qty}}", "{{item_amount}}",
    ]
    for i, tk in enumerate(tokens):
        cell = items.rows[1].cells[i]
        cell.width = widths[i]
        set_cell_text(cell, tk, size=9, align=WD_ALIGN_PARAGRAPH.RIGHT if i in (4, 5) else None)

    # Charges
    section_title(doc, "Installation, Transport, AMC & Warranty")
    cw = [Mm(50), Mm(86), Mm(40)]
    charges = doc.add_table(rows=2, cols=3)
    charges.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_borders(charges)
    header_row(charges, ["Item", "Details", "Amount"], cw, right_cols=(2,))
    for i, tk in enumerate(["{{charge_label}}", "{{charge_details}}", "{{charge_amount}}"]):
        cell = charges.rows[1].cells[i]
        cell.width = cw[i]
        set_cell_text(cell, tk, size=9, align=WD_ALIGN_PARAGRAPH.RIGHT if i == 2 else None)

    # Totals
    section_title(doc, "Summary")
    totals = doc.add_table(rows=4, cols=2)
    totals.alignment = WD_TABLE_ALIGNMENT.RIGHT
    cell_borders(totals)
    rows = [
        ("Subtotal", "{{subtotal}}", False),
        ("Discount ({{discount_pct}}%)", "{{discount_amount}}", False),
        ("Tax ({{tax_rate}})", "{{tax_amount}}", False),
        ("Grand Total", "{{grand_total}}", True),
    ]
    for i, (label, value, strong) in enumerate(rows):
        a, b = totals.rows[i].cells
        a.width = Mm(60)
        b.width = Mm(40)
        set_cell_text(a, label, size=11 if strong else 9.5, bold=strong, color=NAVY if strong else BODY)
        set_cell_text(
            b, value, size=11 if strong else 9.5, bold=strong,
            color=NAVY if strong else BODY, align=WD_ALIGN_PARAGRAPH.RIGHT,
        )
        if strong:
            shade(a, "F6F3EC")
            shade(b, "F6F3EC")

    # Commercial terms
    section_title(doc, "Commercial Terms")
    para(doc, "Payment terms: {{payment_terms}}", size=9.5)
    para(doc, "Warranty: {{warranty_months}} months. {{warranty_terms}}", size=9.5)
    para(doc, "AMC: {{amc_tier}} {{amc_price}}", size=9.5)
    para(doc, "{{commercial_notes}}", size=9.5)

    # Next steps
    section_title(doc, "Next Steps")
    para(doc, "{{next_steps}}", size=10)

    # Footer
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer.add_run("{{quotation_number}} | {{company_name}}")
    fr.font.name = "Arial"
    fr.font.size = Pt(8)
    fr.font.color.rgb = GREY

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    doc.save(OUTPUT_FILE)


if __name__ == "__main__":
    build()
