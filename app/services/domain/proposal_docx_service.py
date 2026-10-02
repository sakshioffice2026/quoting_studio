from __future__ import annotations

import copy
import io
import re

from docx import Document

TOKEN_RE = re.compile(r"\{\{\s*([a-zA-Z0-9_]+)\s*\}\}")

ITEM_ANCHOR = "{{item_label}}"
CHARGE_ANCHOR = "{{charge_label}}"
DISCOUNT_ANCHOR = "{{discount_amount}}"


# ------------------------------------------------------------------ #
#  Text helpers
# ------------------------------------------------------------------ #

def _as_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (list, tuple)):
        return "\n".join(f"{i}. {v}" for i, v in enumerate(value, start=1))
    return str(value)


def _set_run_text(run, text: str) -> None:
    parts = text.split("\n")
    run.text = parts[0]
    for part in parts[1:]:
        run.add_break()
        run.add_text(part)


def _replace_in_paragraph(paragraph, mapping: dict) -> None:
    full = paragraph.text
    if "{{" not in full:
        return
    runs = paragraph.runs
    if not runs:
        return

    def _sub(match):
        return _as_text(mapping.get(match.group(1), ""))

    new_text = TOKEN_RE.sub(_sub, full)
    if new_text == full:
        return

    _set_run_text(runs[0], new_text)
    for extra in runs[1:]:
        extra.text = ""


def _iter_tables(container):
    for table in container.tables:
        yield table
        for row in table.rows:
            for cell in row.cells:
                yield from _iter_tables(cell)


def _iter_paragraphs(container):
    for p in container.paragraphs:
        yield p
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from _iter_paragraphs(cell)


def _replace_everywhere(doc, mapping: dict) -> None:
    for p in _iter_paragraphs(doc):
        _replace_in_paragraph(p, mapping)

    for section in doc.sections:
        for part in (section.header, section.footer,
                     section.first_page_header, section.first_page_footer):
            if part is None:
                continue
            for p in _iter_paragraphs(part):
                _replace_in_paragraph(p, mapping)


# ------------------------------------------------------------------ #
#  Table helpers
# ------------------------------------------------------------------ #

def _row_text(row) -> str:
    return " ".join(cell.text for cell in row.cells)


def _find_template_row(doc, anchor: str):
    for table in _iter_tables(doc):
        for row in table.rows:
            if anchor in _row_text(row):
                return table, row
    return None, None


def _remove_row(table, row) -> None:
    table._tbl.remove(row._tr)


def _fill_repeating_rows(doc, anchor: str, rows: list[dict]) -> None:
    table, template_row = _find_template_row(doc, anchor)
    if table is None:
        return

    if not rows:
        _remove_row(table, template_row)
        return

    template_xml = copy.deepcopy(template_row._tr)
    anchor_tr = template_row._tr

    for row_data in rows:
        new_tr = copy.deepcopy(template_xml)
        anchor_tr.addprevious(new_tr)

    table._tbl.remove(anchor_tr)

    # fill the freshly inserted rows
    start = len(table.rows) - len(rows)
    for offset, row_data in enumerate(rows):
        row = table.rows[start + offset]
        for cell in row.cells:
            for p in cell.paragraphs:
                _replace_in_paragraph(p, row_data)


def _drop_discount_row_if_empty(doc, scalars: dict) -> None:
    pct = str(scalars.get("discount_pct", "")).strip()
    if pct not in ("", "0"):
        return
    table, row = _find_template_row(doc, DISCOUNT_ANCHOR)
    if table is not None:
        _remove_row(table, row)


# ------------------------------------------------------------------ #
#  Public API
# ------------------------------------------------------------------ #

def build_mapping(context: dict, copy_data: dict) -> dict:
    mapping = dict(context["scalars"])
    mapping["executive_summary"] = _as_text(copy_data.get("executive_summary"))
    mapping["scope_of_work"] = _as_text(copy_data.get("scope_of_work"))
    mapping["commercial_notes"] = _as_text(copy_data.get("commercial_notes"))
    mapping["next_steps"] = _as_text(copy_data.get("next_steps"))
    return mapping


def fill_docx_template(template_path: str, context: dict, copy_data: dict) -> bytes:
    """
    template_path : path to tenant/default .docx containing {{tokens}}
    context       : output of proposal_context_service.build_proposal_context
    copy_data     : dict with executive_summary, scope_of_work, commercial_notes, next_steps
    returns       : filled .docx as bytes
    """
    doc = Document(template_path)
    mapping = build_mapping(context, copy_data)

    _fill_repeating_rows(doc, ITEM_ANCHOR, context["item_rows"])
    _fill_repeating_rows(doc, CHARGE_ANCHOR, context["charge_rows"])
    _drop_discount_row_if_empty(doc, context["scalars"])

    _replace_everywhere(doc, mapping)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
