from __future__ import annotations

import io
import json
import os
from datetime import datetime

from docx import Document
from flask import current_app
from werkzeug.utils import secure_filename

TEMPLATE_DIR = "proposal_templates"
DEFAULT_MARKER = "_default.json"
BUILTIN_DEFAULT = "default_proposal.docx"
MAX_TEMPLATE_BYTES = 5 * 1024 * 1024
REQUIRED_TOKENS = ("{{item_label}}", "{{grand_total}}")


# ------------------------------------------------------------------ #
#  Paths
# ------------------------------------------------------------------ #

def tenant_dir(tenant_id: int) -> str:
    path = os.path.join(current_app.config["UPLOAD_FOLDER"], TEMPLATE_DIR, str(tenant_id))
    os.makedirs(path, exist_ok=True)
    return path


def builtin_template_path() -> str:
    return os.path.join(current_app.root_path, "proposal_templates", BUILTIN_DEFAULT)


def _marker_path(tenant_id: int) -> str:
    return os.path.join(tenant_dir(tenant_id), DEFAULT_MARKER)


# ------------------------------------------------------------------ #
#  Default selection
# ------------------------------------------------------------------ #

def get_default_name(tenant_id: int) -> str | None:
    path = _marker_path(tenant_id)
    if not os.path.isfile(path):
        return None
    try:
        with open(path, "r", encoding="utf-8") as fh:
            name = json.load(fh).get("default")
    except (OSError, ValueError):
        return None
    if name and os.path.isfile(os.path.join(tenant_dir(tenant_id), name)):
        return name
    return None


def set_default(tenant_id: int, filename: str | None) -> None:
    path = _marker_path(tenant_id)
    if not filename:
        if os.path.isfile(path):
            os.remove(path)
        return

    safe = secure_filename(filename)
    if not os.path.isfile(os.path.join(tenant_dir(tenant_id), safe)):
        raise FileNotFoundError("Template not found")

    with open(path, "w", encoding="utf-8") as fh:
        json.dump({"default": safe}, fh)


# ------------------------------------------------------------------ #
#  List / save / delete
# ------------------------------------------------------------------ #

def list_templates(tenant_id: int) -> list[dict]:
    folder = tenant_dir(tenant_id)
    default_name = get_default_name(tenant_id)
    rows = []

    for name in sorted(os.listdir(folder)):
        if not name.lower().endswith(".docx"):
            continue
        full = os.path.join(folder, name)
        stat = os.stat(full)
        rows.append(
            {
                "filename": name,
                "size_kb": round(stat.st_size / 1024, 1),
                "modified": datetime.fromtimestamp(stat.st_mtime),
                "is_default": name == default_name,
            }
        )
    return rows


def _unique_name(folder: str, name: str) -> str:
    base, ext = os.path.splitext(name)
    candidate = name
    counter = 2
    while os.path.exists(os.path.join(folder, candidate)):
        candidate = f"{base}_{counter}{ext}"
        counter += 1
    return candidate


def save_template(tenant_id: int, file_storage) -> str:
    """Validates and stores an uploaded .docx. Returns the stored filename."""
    original = secure_filename(file_storage.filename or "")
    if not original.lower().endswith(".docx"):
        raise ValueError("Only .docx files are allowed.")

    data = file_storage.read()
    if not data:
        raise ValueError("The file is empty.")
    if len(data) > MAX_TEMPLATE_BYTES:
        raise ValueError("Template is larger than 5 MB.")
    if data[:2] != b"PK":
        raise ValueError("The file is not a valid .docx document.")

    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:
        raise ValueError("The file could not be opened as a Word document.") from exc

    text = _collect_text(doc)
    missing = [t for t in REQUIRED_TOKENS if t not in text]
    if missing:
        raise ValueError("Template is missing required tokens: " + ", ".join(missing))

    folder = tenant_dir(tenant_id)
    stored = _unique_name(folder, original)
    with open(os.path.join(folder, stored), "wb") as fh:
        fh.write(data)
    return stored


def _collect_text(doc) -> str:
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def delete_template(tenant_id: int, filename: str) -> None:
    safe = secure_filename(filename)
    full = os.path.join(tenant_dir(tenant_id), safe)
    if not os.path.isfile(full):
        raise FileNotFoundError("Template not found")

    if get_default_name(tenant_id) == safe:
        set_default(tenant_id, None)
    os.remove(full)


# ------------------------------------------------------------------ #
#  Resolve
# ------------------------------------------------------------------ #

def resolve_template_path(tenant_id: int, template_filename: str | None = None) -> str:
    folder = tenant_dir(tenant_id)

    if template_filename:
        candidate = os.path.join(folder, secure_filename(template_filename))
        if os.path.isfile(candidate):
            return candidate

    default_name = get_default_name(tenant_id)
    if default_name:
        return os.path.join(folder, default_name)

    builtin = builtin_template_path()
    if not os.path.isfile(builtin):
        raise FileNotFoundError(
            "Default proposal template not found. Run scripts/build_default_proposal_template.py"
        )
    return builtin
